"""Generate the VSA Advisor Bot user-evaluation study documents.

This utility creates the formatted Word documents used for the project's
target-user evaluation: a participant information and consent form, and a
moderated usability/evaluation questionnaire. The generated materials cover
consent, privacy, task completion, Likert-scale feedback, understanding checks,
open comments and moderator observations.

The script supports the project's usability and comprehension evaluation only.
It does not run the VSA strategy, train or score the machine-learning models,
or interact with MT5 trading logic. Its role is to produce consistent study
documents that can be used to collect structured feedback from target users.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(__file__).resolve().parents[1]
OUT = PROJECT / "docs" / "user_study"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)
LIGHT_FILL = "F2F4F7"
PALE_BLUE = "E8EEF5"
INK = RGBColor(25, 30, 36)


def set_run_font(run, size=11, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def configure_document(doc: Document, running_label: str):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run(running_label), size=9, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    set_run_font(fp.add_run("VSA Advisor Bot  |  Page "), size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)


def add_title_block(doc: Document, title: str, subtitle: str, label: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(p.add_run(label.upper()), size=9, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(title), size=23, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run(subtitle), size=12.5, color=MUTED)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total_dxa = sum(int(round(width * 1440)) for width in widths_inches)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_inches:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(round(width * 1440))))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            dxa = int(round(widths_inches[idx] * 1440))
            cell.width = Inches(widths_inches[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(dxa))
            tc_w.set(qn("w:type"), "dxa")


def add_metadata(doc: Document, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(rows):
        left, right = table.rows[i].cells
        set_cell_shading(left, LIGHT_FILL)
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label), bold=True, color=DARK_BLUE)
        p = right.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value))
    set_table_geometry(table, [1.55, 4.95])


def add_note_box(doc: Document, label: str, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(f"{label}: "), bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(text))
    set_table_geometry(table, [6.5])


def add_checkbox(doc: Document, text: str, bold_prefix: str | None = None, space_after=7):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.first_line_indent = Inches(-0.02)
    p.paragraph_format.space_after = Pt(space_after)
    set_run_font(p.add_run("☐  "), size=12, color=DARK_BLUE)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))


def add_signature_line(doc: Document, label: str, width_chars=58):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(f"{label}: "), bold=True)
    set_run_font(p.add_run("_" * width_chars), color=MUTED)


def add_response_lines(doc: Document, count=3):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "B8BEC7")
        p_bdr.append(bottom)
        p_pr.append(p_bdr)
        set_run_font(p.add_run(" "))


def create_consent_form(path: Path):
    doc = Document()
    configure_document(doc, "Participant Information & Consent Form")
    add_title_block(
        doc,
        "VSA Advisor Bot User Evaluation",
        "Participant information and signed consent form",
        "Final Year Project • Financial Advisor Bot Template 4.2",
    )
    add_metadata(doc, [
        ("Researcher", "Syed Abdullah Humayun"),
        ("Institution / programme", "________________________________________________"),
        ("Supervisor", "________________________________________________"),
        ("Study duration", "Approximately 15–20 minutes"),
        ("Target participants", "Adults interested in trading, finance or computing; novice users are especially relevant"),
    ])

    doc.add_heading("Why am I being invited?", level=1)
    doc.add_paragraph(
        "You are invited to evaluate the VSA Advisor Bot, an educational prototype combining a defined Volume Spread "
        "Analysis (VSA) strategy with a trained Logistic Regression model. Your feedback will assess whether its interface, "
        "explanations and demo recommendations are understandable to non-technical users."
    )

    doc.add_heading("What will I be asked to do?", level=1)
    for text in [
        "Complete short navigation and interpretation tasks on the Streamlit website.",
        "Review one controlled VSA + ML recommendation and explain its meaning.",
        "Review the news warning, risk percentage and managed-exit explanation.",
        "Rate usability, clarity, usefulness and trust; optional comments are welcome.",
    ]:
        add_checkbox(doc, text, space_after=3)

    doc.add_heading("Risk, privacy and voluntary participation", level=1)
    doc.add_paragraph(
        "You will not deposit money, disclose credentials, place a trade or provide financial information. The prototype is "
        "educational, uses demonstration material and does not guarantee profit. Participation is voluntary: you may skip a "
        "question or stop before submitting the questionnaire without giving a reason."
    )
    doc.add_paragraph(
        "The signed consent sheet will be stored separately from coded questionnaire responses. Results will be reported in "
        "aggregate; short anonymous quotations may be used. Names, signatures and identifying information will not be published."
    )
    add_note_box(
        doc,
        "Important",
        "Do not use the prototype to make a real-money trading decision. Historical performance and model probability do not establish future profitability.",
    )

    consent_heading = doc.add_heading("Consent declaration", level=1)
    consent_heading.paragraph_format.page_break_before = True
    doc.add_paragraph("Please tick every box before signing. Ask the researcher if any statement is unclear.")
    statements = [
        "I confirm that I am 18 years of age or older.",
        "I have read and understood the participant information on the previous page.",
        "I understand that participation is voluntary and that I may stop before submitting my responses.",
        "I understand that this is an educational/demo system and not personal financial advice.",
        "I understand what information will be collected and that the signed form will be stored separately from coded responses.",
        "I understand that anonymised aggregate results and short anonymous quotations may appear in the final project report.",
        "I understand that I must not enter broker credentials, financial details or other sensitive information.",
        "I freely agree to participate in this user evaluation.",
    ]
    for statement in statements:
        add_checkbox(doc, statement)

    add_signature_line(doc, "Participant full name")
    add_signature_line(doc, "Participant signature")
    add_signature_line(doc, "Date", 24)
    add_signature_line(doc, "Participant code assigned by researcher", 29)
    add_signature_line(doc, "Researcher signature")
    add_signature_line(doc, "Date", 24)

    add_note_box(
        doc,
        "Researcher check",
        "Give the participant time to ask questions. Do not begin the questionnaire unless all consent statements are ticked and the form is signed.",
    )
    doc.save(path)


def add_table_header(table, labels):
    row = table.rows[0]
    set_repeat_table_header(row)
    for index, label in enumerate(labels):
        set_cell_shading(row.cells[index], PALE_BLUE)
        p = row.cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index != 1 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(label), size=9.5, bold=True, color=DARK_BLUE)


def create_questionnaire(path: Path):
    doc = Document()
    configure_document(doc, "Target-User Evaluation Questionnaire")
    add_title_block(
        doc,
        "VSA Advisor Bot Evaluation",
        "Target-user task sheet and questionnaire",
        "Anonymous response form • Complete after signing the separate consent form",
    )
    add_metadata(doc, [
        ("Participant code", "______________________"),
        ("Session date", "______________________"),
        ("Researcher / moderator", "______________________"),
        ("Expected duration", "15–20 minutes"),
    ])
    add_note_box(
        doc,
        "Before starting",
        "Use only the demonstration website and demo MT5 evidence. Do not enter account credentials or place a real-money trade.",
    )

    doc.add_heading("A. Participant background", level=1)
    doc.add_paragraph("Tick the option that best describes you. These questions do not request financial information.")
    add_checkbox(doc, "I confirm that I am 18 years of age or older.")
    doc.add_paragraph("Primary background:")
    for option in [
        "Novice or interested trader", "Finance or business student", "Computing student",
        "Experienced trader", "General user", "Other: ______________________________",
    ]:
        add_checkbox(doc, option)
    doc.add_paragraph("Previous experience with trading platforms:")
    for option in ["None", "Less than 1 year", "1–3 years", "More than 3 years"]:
        add_checkbox(doc, option)

    doc.add_heading("B. Moderated usability tasks", level=1)
    doc.add_paragraph(
        "Complete each task without coaching where possible. The moderator records the outcome; the participant may comment aloud."
    )
    tasks = [
        ("1", "Open the ML Advisor page and identify the selected model."),
        ("2", "State the direction, scenario number, LR probability and final decision shown."),
        ("3", "Explain in your own words what the LR probability means—and what it does not guarantee."),
        ("4", "Find the VSA + ML Algorithm page and identify how Scenario 1, 2 or 3 is confirmed."),
        ("5", "Find the news-risk status and explain why a raw ML candidate may become WATCH_ONLY."),
        ("6", "Find the risk-management explanation: scenario risk, 70% at 1R, breakeven, and 30% runner to 4R."),
        ("7", "Find the LR versus RF evidence and state why LR is deployed in this prototype."),
        ("8", "Find the MT5 integration evidence and identify whether a demo trade was executed."),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    add_table_header(table, ["#", "Task", "No help", "Some help", "Not completed"])
    for number, task in tasks:
        cells = table.add_row().cells
        values = [number, task, "☐", "☐", "☐"]
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=9.5 if i == 1 else 11)
    set_table_geometry(table, [0.35, 3.75, 0.80, 0.80, 0.80])

    doc.add_page_break()
    doc.add_heading("C. Five-point evaluation", level=1)
    doc.add_paragraph("Circle or tick one response for every statement: 1 = strongly disagree, 5 = strongly agree.")
    statements = [
        "The purpose of the VSA Advisor Bot was clear.",
        "The website was easy to navigate.",
        "The ML Advisor page opened quickly enough for comfortable use.",
        "The BUY/SELL direction and WATCH_ONLY/TRADE_CANDIDATE decision were easy to identify.",
        "The three VSA scenarios were explained clearly.",
        "I understood that VSA generates a candidate before ML evaluates it.",
        "I understood that the LR probability estimates setup success rather than an exact future price.",
        "The LR explanation helped me understand why the model supported or rejected a setup.",
        "The comparison between Logistic Regression and Random Forest was understandable.",
        "The high-impact USD news warning was clear.",
        "The 0.5%/1% risk rules and scenario-specific stop were clear.",
        "The 70% at 1R, breakeven and 30% runner to 4R process was clear.",
        "The warnings prevented me from assuming that profit is guaranteed.",
        "I would consider using this prototype for education or demo trading only.",
        "Overall, the prototype was useful for understanding VSA + ML decision support.",
    ]
    likert = doc.add_table(rows=1, cols=7)
    likert.style = "Table Grid"
    add_table_header(likert, ["#", "Statement", "1", "2", "3", "4", "5"])
    for idx, statement in enumerate(statements, start=1):
        cells = likert.add_row().cells
        values = [str(idx), statement, "☐", "☐", "☐", "☐", "☐"]
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run(value), size=9.2 if i == 1 else 10)
    set_table_geometry(likert, [0.35, 4.15, 0.40, 0.40, 0.40, 0.40, 0.40])

    doc.add_page_break()
    doc.add_heading("D. Understanding checks", level=1)
    questions = [
        ("1. What does the LR probability represent?", [
            "A guaranteed chance of profit", "The estimated chance that this formal VSA setup reaches 1R before its stop",
            "The exact future price of gold", "The percentage of account capital to risk",
        ]),
        ("2. Why can a TRADE_CANDIDATE become WATCH_ONLY?", [
            "The system may apply a high-impact-news or safety block", "Random Forest automatically replaces LR",
            "The system guarantees a loss", "MT5 changes the strategy rules",
        ]),
        ("3. Which model is deployed in the current prototype?", [
            "Logistic Regression", "Random Forest", "A language model", "No trained model",
        ]),
    ]
    for question, answers in questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(question), bold=True)
        for answer in answers:
            add_checkbox(doc, answer)

    doc.add_page_break()
    doc.add_heading("E. Open feedback", level=1)
    prompts = [
        "What was the clearest or most useful part of the prototype?",
        "What was confusing, slow or difficult to find?",
        "Did any wording make the recommendation appear more certain than it really is? Explain.",
        "What should be improved before another novice user tests the system?",
        "Any other comments?",
    ]
    for prompt in prompts:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(prompt), bold=True, color=DARK_BLUE)
        add_response_lines(doc, 2)

    doc.add_heading("F. Moderator observations", level=1)
    add_signature_line(doc, "Total session time", 24)
    add_signature_line(doc, "Tasks completed without help", 20)
    add_signature_line(doc, "Tasks completed with some help", 17)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Observed usability problems or notable comments:"), bold=True, color=DARK_BLUE)
    add_response_lines(doc, 3)
    add_note_box(
        doc,
        "Data handling",
        "Enter only the participant code in the analysis dataset. Keep this questionnaire and the signed consent form secure and separate. Report aggregated results and anonymised quotations only.",
    )
    doc.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    create_consent_form(OUT / "VSA_Advisor_Participant_Information_and_Consent_Form.docx")
    create_questionnaire(OUT / "VSA_Advisor_Target_User_Evaluation_Questionnaire.docx")
    print(f"Created study documents in {OUT}")


if __name__ == "__main__":
    main()
